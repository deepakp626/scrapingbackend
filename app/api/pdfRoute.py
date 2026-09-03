from fastapi import Form, APIRouter, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse, FileResponse
from fastapi.responses import Response
from typing import Optional
from pypdf import PdfReader
import zipfile


from data_pipeline.pdf_data_pytesseract import PDFTextExtractor

import io
import os
import shutil
import subprocess
import tempfile
from pathlib import Path
import pikepdf

import pymupdf
from PIL import Image
from docx import Document
from pdfminer.high_level import extract_text_to_fp
from pdfminer.layout import LAParams


router = APIRouter(
    prefix="/pdf",
    tags=["PDF Tools"]
)


# ============================================================
# PDF COMPRESS
# ============================================================

@router.post("/compress")
async def compress_pdf(
    file: UploadFile = File(...),
    percentage: float = 50.0
):
    """
    Compress PDF using PyMuPDF and Pillow.

    POST /pdf/compress
    """

    if file.content_type != "application/pdf":
        raise HTTPException(
            status_code=400,
            detail="Only PDF files are allowed."
        )

    try:
        pdf_bytes = await file.read()

        if not pdf_bytes:
            raise HTTPException(
                status_code=400,
                detail="Uploaded PDF is empty."
            )

        # Clamp percentage
        compress_pct = max(
            1.0,
            min(100.0, float(percentage))
        )

        # Open PDF
        input_pdf = pymupdf.open(
            stream=pdf_bytes,
            filetype="pdf"
        )

        if input_pdf.page_count == 0:
            input_pdf.close()

            raise HTTPException(
                status_code=400,
                detail="PDF contains no pages."
            )

        # Create output PDF
        output_pdf = pymupdf.open()

        output_pdf.insert_pdf(input_pdf)

        # Compression settings
        scale = max(
            0.3,
            min(
                1.0,
                1.0 - (compress_pct / 100.0) * 0.6
            )
        )

        quality = int(
            max(
                15,
                min(
                    95,
                    100 - (compress_pct * 0.8)
                )
            )
        )

        processed_xrefs = set()

        for page in output_pdf:

            for img_info in page.get_images(full=True):

                xref = img_info[0]

                if xref in processed_xrefs:
                    continue

                processed_xrefs.add(xref)

                try:
                    base_image = output_pdf.extract_image(xref)

                    if not base_image:
                        continue

                    raw_img_bytes = base_image.get("image")

                    if not raw_img_bytes:
                        continue

                    pil_img = Image.open(
                        io.BytesIO(raw_img_bytes)
                    )

                    # Resize
                    if scale < 1.0:

                        new_width = max(
                            1,
                            int(pil_img.width * scale)
                        )

                        new_height = max(
                            1,
                            int(pil_img.height * scale)
                        )

                        pil_img = pil_img.resize(
                            (new_width, new_height),
                            Image.Resampling.LANCZOS
                        )

                    # Convert image to RGB
                    if pil_img.mode in ("RGBA", "P"):

                        background = Image.new(
                            "RGB",
                            pil_img.size,
                            (255, 255, 255)
                        )

                        if pil_img.mode == "RGBA":

                            background.paste(
                                pil_img,
                                mask=pil_img.split()[-1]
                            )

                        else:

                            background.paste(pil_img)

                        pil_img = background

                    elif pil_img.mode != "RGB":

                        pil_img = pil_img.convert("RGB")

                    # Compress image
                    out_img_io = io.BytesIO()

                    pil_img.save(
                        out_img_io,
                        format="JPEG",
                        quality=quality,
                        optimize=True
                    )

                    new_image_bytes = out_img_io.getvalue()

                    # Replace only if smaller
                    if len(new_image_bytes) < len(raw_img_bytes):

                        page.replace_image(
                            xref,
                            stream=new_image_bytes
                        )

                except Exception:
                    # Ignore individual image errors
                    continue

        # Save compressed PDF
        compressed_bytes = io.BytesIO()

        output_pdf.save(
            compressed_bytes,
            garbage=4,
            deflate=True,
            clean=True,
            use_objstms=True
        )

        output_pdf.close()
        input_pdf.close()

        compressed_bytes.seek(0)

        compressed_size = len(
            compressed_bytes.getvalue()
        )

        return StreamingResponse(
            compressed_bytes,
            media_type="application/pdf",
            headers={
                "Content-Disposition":
                    'attachment; filename="compressed.pdf"',

                "X-Original-Size":
                    str(len(pdf_bytes)),

                "X-Compressed-Size":
                    str(compressed_size),

                "X-Compression-Percentage":
                    str(compress_pct)
            }
        )

    except HTTPException:
        raise

    except Exception as e:

        raise HTTPException(
            status_code=500,
            detail=f"Failed to compress PDF: {str(e)}"
        )


# ============================================================
# PDF TO WORD
# ============================================================

@router.post("/to-word")
async def pdf_to_word(
    file: UploadFile = File(...)
):
    """
    Convert PDF to Word.

    PyMuPDF:
        Extracts PDF text.

    python-docx:
        Creates DOCX.

    POST /pdf/to-word
    """

    if file.content_type != "application/pdf":
        raise HTTPException(
            status_code=400,
            detail="Only PDF files are allowed."
        )

    try:

        pdf_bytes = await file.read()

        if not pdf_bytes:

            raise HTTPException(
                status_code=400,
                detail="Uploaded PDF is empty."
            )

        # Open PDF
        pdf_document = pymupdf.open(
            stream=pdf_bytes,
            filetype="pdf"
        )

        if pdf_document.page_count == 0:

            pdf_document.close()

            raise HTTPException(
                status_code=400,
                detail="PDF contains no pages."
            )

        # Create Word
        document = Document()

        # Process pages
        for page_number, page in enumerate(
            pdf_document
        ):

            text = page.get_text("text")

            document.add_heading(
                f"Page {page_number + 1}",
                level=2
            )

            if text.strip():

                lines = text.splitlines()

                for line in lines:

                    if line.strip():

                        document.add_paragraph(
                            line
                        )

            else:

                document.add_paragraph(
                    "[No extractable text found on this page]"
                )

            # Page break
            if page_number < pdf_document.page_count - 1:

                document.add_page_break()

        pdf_document.close()

        # Save DOCX in memory
        word_file = io.BytesIO()

        document.save(word_file)

        word_file.seek(0)

        return StreamingResponse(
            word_file,
            media_type=(
                "application/vnd.openxmlformats-"
                "officedocument.wordprocessingml.document"
            ),
            headers={
                "Content-Disposition":
                    'attachment; filename="converted.docx"'
            }
        )

    except HTTPException:
        raise

    except Exception as e:

        raise HTTPException(
            status_code=500,
            detail=f"Failed to convert PDF to Word: {str(e)}"
        )



# ============================================================
# PDF LOCK
# ============================================================

@router.post("/lock-pdf", summary="Password-protect an uploaded PDF file")
async def lock_pdf(
    file: UploadFile = File(..., description="The PDF file to lock"),
    password: str = Form(..., description="The password to encrypt the PDF with"),
):
    # 1. Validate file extension/type
    if not file.filename.lower().endswith(".pdf"):
        raise HTTPException(
            status_code=400, detail="Invalid file format. Only PDF files are allowed."
        )

    temp_input_path = None
    temp_output_path = None

    # Use temporary files to securely handle incoming stream and processing with pikepdf
    try:
        with tempfile.NamedTemporaryFile(
            delete=False, suffix=".pdf"
        ) as temp_input, tempfile.NamedTemporaryFile(
            delete=False, suffix=".pdf"
        ) as temp_output:

            # Save uploaded content into the temporary input file
            shutil.copyfileobj(file.file, temp_input)
            temp_input_path = temp_input.name
            temp_output_path = temp_output.name

        # 2. Open and encrypt the PDF using pikepdf
        try:
            with pikepdf.open(temp_input_path) as pdf:
                # Configure 256-bit AES encryption
                encryption = pikepdf.Encryption(
                    owner=password,  # Owner password has full control
                    user=password,   # User password required to open the PDF
                    R=6,             # R=6 specifies modern AES-256 encryption standard
                )
                # Save the encrypted file to the output path
                pdf.save(temp_output_path, encryption=encryption)
                
        except pikepdf.PasswordError:
            # Catches cases where the uploaded PDF is already encrypted/locked
            raise HTTPException(
                status_code=400,
                detail="The uploaded PDF file is already password-protected. Please upload an unlocked PDF."
            )

        # 3. Read the encrypted file into memory to stream back to the client
        with open(temp_output_path, "rb") as locked_file:
            pdf_bytes = locked_file.read()

        # Generate a clean output filename
        output_filename = f"locked_{file.filename}"

        # 4. Return as a stream response
        return StreamingResponse(
            io.BytesIO(pdf_bytes),
            media_type="application/pdf",
            headers={
                "Content-Disposition": f'attachment; filename="{output_filename}"'
            },
        )

    except HTTPException as he:
        # Re-raise explicit HTTP exceptions (like our PasswordError catch)
        raise he
    except pikepdf.PdfError as e:
        raise HTTPException(
            status_code=400, detail=f"Failed to process PDF file: {str(e)}"
        )
    except Exception as e:
        raise HTTPException(
            status_code=500, detail=f"An unexpected error occurred: {str(e)}"
        )

    finally:
        # Cleanup uploaded file stream and temporary files safely
        file.file.close()
        for path in [temp_input_path, temp_output_path]:
            if path and os.path.exists(path):
                try:
                    os.remove(path)
                except Exception:
                    pass


# ============================================================
# PDF UNLOCK
# ============================================================

@router.post("/unlock-pdf", summary="Unlock or remove restrictions from a PDF file")
async def unlock_pdf(
    file: UploadFile = File(..., description="The PDF file to unlock"),
    password: Optional[str] = Form(default="", description="The password to unlock the PDF (if protected)"),
):
    # 1. Validate file extension/type
    if not file.filename.lower().endswith(".pdf"):
        raise HTTPException(
            status_code=400, detail="Invalid file format. Only PDF files are allowed."
        )

    temp_input_path = None
    temp_output_path = None

    try:
        # Create temporary files for handling input and output streams
        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as temp_input:
            shutil.copyfileobj(file.file, temp_input)
            temp_input_path = temp_input.name

        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as temp_output:
            temp_output_path = temp_output.name

        # 2. Open the PDF using the provided password or empty password ("")
        pwd = password if password is not None else ""
        try:
            with pikepdf.open(temp_input_path, password=pwd) as pdf:
                # Saving without encryption parameters strips all security/restrictions
                pdf.save(temp_output_path)
                
        except pikepdf.PasswordError:
            if pwd:
                raise HTTPException(
                    status_code=400,
                    detail="Incorrect password. Failed to unlock the PDF."
                )
            else:
                raise HTTPException(
                    status_code=400,
                    detail="This PDF is password-protected. Please provide the password to unlock it."
                )

        # 3. Read the unlocked file into memory
        with open(temp_output_path, "rb") as unlocked_file:
            pdf_bytes = unlocked_file.read()

        output_filename = f"unlocked_{file.filename}"

        # 4. Return as a stream response
        return StreamingResponse(
            io.BytesIO(pdf_bytes),
            media_type="application/pdf",
            headers={
                "Content-Disposition": f'attachment; filename="{output_filename}"'
            },
        )

    except HTTPException as he:
        raise he
    except pikepdf.PdfError as e:
        raise HTTPException(
            status_code=400, detail=f"Failed to process PDF file: {str(e)}"
        )
    except Exception as e:
        raise HTTPException(
            status_code=500, detail=f"An unexpected error occurred: {str(e)}"
        )

    finally:
        # Cleanup uploaded file stream and temporary files safely
        file.file.close()
        for path in [temp_input_path, temp_output_path]:
            if path and os.path.exists(path):
                try:
                    os.remove(path)
                except Exception:
                    pass



# Initialize PDF text extractor
pdf_extractor = PDFTextExtractor()

@router.post("/ocr-pdf", summary="Extract and OCR text from a PDF document")
async def ocr_pdf(file: UploadFile = File(..., description="The PDF file to extract text from")):
    """
    Extract text from uploaded PDF using native extraction with intelligent OCR fallback.
    Returns full extracted text, per-page breakdown, word & character counts, and performance metrics.

    POST /pdf/ocr-pdf
    """
    # 1. Validate file format
    if not file.filename or not file.filename.lower().endswith(".pdf"):
        raise HTTPException(
            status_code=400,
            detail="Invalid file format. Only PDF files (.pdf) are allowed."
        )

    try:
        content = await file.read()
        
        if not content:
            raise HTTPException(
                status_code=400,
                detail="Uploaded PDF file is empty."
            )

        extraction_result = await pdf_extractor.extract_pdf_text(content)

        extracted_text = extraction_result.get("text", "")
        total_pages = extraction_result.get("total_pages", 0)
        word_count = extraction_result.get("word_count", 0)
        char_count = extraction_result.get("char_count", 0)
        pages = extraction_result.get("pages", [])
        stats = extraction_result.get("stats", {
            "normal_pdf_time": extraction_result.get("normal_pdf_time", 0),
            "ocr_time": extraction_result.get("ocr_time", 0),
            "total_time": extraction_result.get("total_time", 0),
            "ocr_used": extraction_result.get("ocr_used", False)
        })

        # print("data","",{
        #     "success": True,
        #     "message": "File processed successfully",
        #     "filename": file.filename,
        #     "total_pages": total_pages,
        #     "word_count": word_count,
        #     "char_count": char_count,
        #     "extracted_text": extracted_text,
        #     "pages": pages,
        #     "stats": stats,
        #     # Backward compatibility
        #     "CV data": extracted_text
        # })

        return {
            "success": True,
            "message": "File processed successfully",
            "filename": file.filename,
            "total_pages": total_pages,
            "word_count": word_count,
            "char_count": char_count,
            "extracted_text": extracted_text,
            "pages": pages,
            "stats": stats,
            # Backward compatibility
            "CV data": extracted_text
        }

    except HTTPException as he:
        raise he

    except ValueError as ve:
        raise HTTPException(
            status_code=400,
            detail=str(ve)
        )

    except Exception as e:
        error_msg = str(e)

        if "password" in error_msg.lower() or "encrypted" in error_msg.lower():
            raise HTTPException(
                status_code=400,
                detail="This PDF is password-protected. Please unlock it before extracting text."
            )

        if "402" in error_msg or "credits" in error_msg.lower() or "max_tokens" in error_msg.lower():
            raise HTTPException(
                status_code=402,
                detail=f"OpenRouter Credit/Limit Error: {error_msg}. Consider reducing max_tokens or upgrading your plan."
            )

        raise HTTPException(
            status_code=500,
            detail=f"Error processing file: {error_msg}"
        )

    finally:
        await file.close()


# ============================================================
# PDF TO HTML
# ============================================================

@router.post("/to-html", summary="Convert PDF document to HTML")
@router.post("/pdf-to-html", summary="Convert PDF document to HTML (alias)")
async def pdf_to_html(
    file: UploadFile = File(..., description="The PDF file to convert to HTML"),
    engine: str = Form(default="pymupdf", description="Conversion engine: 'pymupdf' or 'pdfminer'"),
):
    """
    Convert an uploaded PDF document into standalone HTML using PyMuPDF and pdfminer.six.

    - PyMuPDF: Preserves styling, typography, exact text placement, and inline graphics.
    - pdfminer.six: Extracts structural text hierarchy and semantic layout.

    POST /pdf/to-html
    POST /pdf/pdf-to-html
    """
    if not file.filename or not file.filename.lower().endswith(".pdf"):
        raise HTTPException(
            status_code=400,
            detail="Invalid file format. Only PDF files (.pdf) are allowed."
        )

    try:
        pdf_bytes = await file.read()
        if not pdf_bytes:
            raise HTTPException(
                status_code=400,
                detail="Uploaded PDF file is empty."
            )

        html_content = ""
        total_pages = 0
        used_engine = engine.lower().strip()

        # Engine 1: pdfminer.six
        if used_engine == "pdfminer":
            try:
                out_fp = io.BytesIO()
                extract_text_to_fp(
                    io.BytesIO(pdf_bytes),
                    out_fp,
                    laparams=LAParams(),
                    output_type="html",
                    codec="utf-8"
                )
                html_content = out_fp.getvalue().decode("utf-8", errors="replace")

                # Retrieve page count with pymupdf for metadata
                try:
                    with pymupdf.open(stream=pdf_bytes, filetype="pdf") as doc:
                        total_pages = doc.page_count
                except Exception:
                    total_pages = 1
            except Exception:
                # Fallback to PyMuPDF on error
                used_engine = "pymupdf"

        # Engine 2: PyMuPDF (Default & fallback)
        if not html_content or used_engine == "pymupdf":
            try:
                doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")
                total_pages = doc.page_count

                if total_pages == 0:
                    doc.close()
                    raise HTTPException(
                        status_code=400,
                        detail="PDF contains no pages."
                    )

                pages_html = []
                for idx, page in enumerate(doc):
                    page_html = page.get_text("html")
                    pages_html.append(f'<div class="pdf-page" id="page-{idx + 1}" data-page="{idx + 1}">\n{page_html}\n</div>')

                doc.close()

                filename_clean = file.filename.rsplit(".", 1)[0]
                html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{filename_clean} - Converted PDF</title>
    <style>
        body {{
            margin: 0;
            padding: 24px;
            background-color: #f8fafc;
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
            color: #1e293b;
            display: flex;
            flex-direction: column;
            align-items: center;
        }}
        .pdf-container {{
            max-width: 900px;
            width: 100%;
            display: flex;
            flex-direction: column;
            gap: 24px;
        }}
        .pdf-page {{
            background: #ffffff;
            border-radius: 8px;
            box-shadow: 0 4px 12px rgba(0, 0, 0, 0.07);
            padding: 32px;
            overflow-x: auto;
            position: relative;
            box-sizing: border-box;
        }}
        .pdf-page img {{
            max-width: 100%;
            height: auto;
        }}
    </style>
</head>
<body>
    <div class="pdf-container">
        {"".join(pages_html)}
    </div>
</body>
</html>"""
            except Exception as me:
                if "password" in str(me).lower() or "encrypted" in str(me).lower():
                    raise HTTPException(
                        status_code=400,
                        detail="This PDF is password-protected. Please unlock it before converting."
                    )
                raise me

        filename_base = file.filename.rsplit(".", 1)[0] if file.filename else "document"
        output_filename = f"{filename_base}.html"

        return StreamingResponse(
            io.BytesIO(html_content.encode("utf-8")),
            media_type="text/html; charset=utf-8",
            headers={
                "Content-Disposition": f'attachment; filename="{output_filename}"',
                "X-Total-Pages": str(total_pages),
                "X-Conversion-Engine": used_engine,
            }
        )

    except HTTPException:
        raise
    except Exception as e:
        error_msg = str(e)
        if "password" in error_msg.lower() or "encrypted" in error_msg.lower():
            raise HTTPException(
                status_code=400,
                detail="This PDF is password-protected. Please unlock it before converting."
            )
        raise HTTPException(
            status_code=500,
            detail=f"Failed to convert PDF to HTML: {error_msg}"
        )
    finally:
        await file.close()



@router.post("/extract-images")
async def extract_images_from_pdf(file: UploadFile = File(...)):
    # Validate file format
    if not file.filename.lower().endswith(".pdf"):
        raise HTTPException(
            status_code=400, detail="Uploaded file must be a PDF."
        )

    try:
        contents = await file.read()
        reader = PdfReader(io.BytesIO(contents))

        zip_buffer = io.BytesIO()
        image_count = 0

        # Create zip archive in memory
        with zipfile.ZipFile(
            zip_buffer, mode="w", compression=zipfile.ZIP_DEFLATED
        ) as zip_file:
            for page_index, page in enumerate(reader.pages, start=1):
                # Extract image objects directly (ignoring text streams)
                for img_index, image_file in enumerate(
                    page.images, start=1
                ):
                    image_count += 1
                    # Format image name: page_1_img_1_filename.png
                    clean_filename = (
                        f"page_{page_index}_img_{img_index}_{image_file.name}"
                    )
                    # Write image binary data directly into the ZIP archive
                    zip_file.writestr(clean_filename, image_file.data)

        # Check if any images were extracted
        if image_count == 0:
            raise HTTPException(
                status_code=404,
                detail="No embedded images were found in the uploaded PDF.",
            )

        zip_filename = f"extracted_images_{file.filename.rsplit('.', 1)[0]}.zip"

        return Response(
            content=zip_buffer.getvalue(),
            media_type="application/zip",
            headers={
                "Content-Disposition": f'attachment; filename="{zip_filename}"'
            },
        )

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=500, detail=f"Error extracting images: {str(e)}"
        )